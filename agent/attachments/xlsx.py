"""大型工作簿逐行读取；完整记录只存放在任务私有目录。"""
import json
import os
import posixpath
import re
import sqlite3
import zipfile
from decimal import Decimal
from lxml import etree
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format, is_timedelta_format
from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.utils.datetime import MAC_EPOCH, WINDOWS_EPOCH, from_excel, from_ISO8601

MAX_XML = 512 * 1024 * 1024
MAX_MEMBER = 256 * 1024 * 1024
MAX_ROWS = 200000
MAX_BOOK_ROWS = 1000000
MAX_CELLS = 20000000


def release(element):
    element.clear()
    while element.getprevious() is not None:
        del element.getparent()[0]


def rich_text(element):
    parts = []
    for child in element:
        if child.tag.endswith("}t"):
            parts.append(child.text or "")
        elif child.tag.endswith("}r"):
            for run in child:
                if run.tag.endswith("}t"):
                    parts.append(run.text or "")
    return "".join(parts)


def numeric_text(value):
    return isinstance(value, str) and re.fullmatch(r"[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?(?:[eE][+-]?\d+)?", value.strip()) is not None


def parse_workbook(source, output, result, block, warning, archive_path):
    with zipfile.ZipFile(source) as archive:
        members = archive.infolist()
        if len(members) > 10000 or sum(item.file_size for item in members) > MAX_XML:
            raise ValueError("Excel 内部 XML 展开大小超过 512 MiB 或结构数量超过限额")
        if any(item.flag_bits & 1 for item in members):
            raise ValueError("Excel 文件已加密，当前无法读取")
        if any(item.file_size > MAX_MEMBER for item in members):
            raise ValueError("Excel 单项内部 XML 超过 256 MiB 解析限额")
        parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
        workbook = etree.fromstring(archive.read("xl/workbook.xml"), parser)
        relations = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"), parser)
        targets = {node.get("Id"): node.get("Target", "") for node in relations if node.get("TargetMode") != "External"}
        properties = workbook.find("{*}workbookPr")
        epoch = MAC_EPOCH if properties is not None and properties.get("date1904") in ["1", "true"] else WINDOWS_EPOCH
        shared = []
        if "xl/sharedStrings.xml" in archive.namelist():
            characters = 0
            with archive.open("xl/sharedStrings.xml") as stream:
                for _, element in etree.iterparse(stream, events=("end",), tag="{*}si", resolve_entities=False, no_network=True, load_dtd=False):
                    text = rich_text(element).replace("x005F_", "")
                    characters += len(text)
                    if len(shared) >= 1000000 or characters > 64 * 1024 * 1024 or len(text) > 1024 * 1024:
                        raise ValueError("Excel 共享文字超过解析限额")
                    shared.append(text)
                    release(element)
        formats = []
        if "xl/styles.xml" in archive.namelist():
            styles = etree.fromstring(archive.read("xl/styles.xml"), parser)
            custom = {int(node.get("numFmtId")): node.get("formatCode", "") for node in styles.findall("{*}numFmts/{*}numFmt")}
            for node in styles.findall("{*}cellXfs/{*}xf"):
                number = int(node.get("numFmtId", "0"))
                code = custom.get(number, BUILTIN_FORMATS.get(number, ""))
                formats.append((is_date_format(code), is_timedelta_format(code)))
        database_path = output / "tables.sqlite"
        connection = sqlite3.connect(database_path)
        os.chmod(database_path, 0o600)
        try:
            connection.execute("PRAGMA cache_size = -8192")
            connection.execute("CREATE TABLE IF NOT EXISTS attachment_rows (table_id TEXT NOT NULL, position INTEGER NOT NULL, row_number INTEGER NOT NULL, cells TEXT NOT NULL, text_mask TEXT NOT NULL, PRIMARY KEY (table_id, position)) WITHOUT ROWID")
            connection.execute("CREATE INDEX IF NOT EXISTS attachment_row_numbers ON attachment_rows (table_id, row_number)")
            budget = dict(rows=0, cells=0)
            for sheet in workbook.findall("{*}sheets/{*}sheet"):
                target = targets.get(sheet.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"), "")
                member = posixpath.normpath(target.lstrip("/") if target.startswith("/") else "xl/" + target)
                if not member.startswith("xl/") or member not in archive.namelist():
                    raise ValueError("Excel 工作表引用无效，不能读取外部位置")
                table_id = "t" + str(len(result["tables"]) + 1)
                # 先前失败成员留下的未发布记录不能进入下一张表。
                connection.execute("DELETE FROM attachment_rows WHERE table_id = ?", (table_id,))
                _parse_sheet(archive, member, sheet, table_id, connection, shared, formats, epoch, budget, result, block, warning, archive_path)
                connection.commit()
        finally:
            connection.close()


def _parse_sheet(archive, member, sheet, table_id, database, shared, formats, epoch, budget, result, block, warning, archive_path):
    name = sheet.get("name", "工作表")
    count = 0
    width = 0
    last_number = 0
    physical_number = 0
    start = None
    first_is_text = False
    merges = []
    gaps = []
    local_warnings = []
    formula_rows = set()

    def note(message, incomplete=True):
        if message not in local_warnings:
            local_warnings.append(message)
            if incomplete:
                warning(name + "：" + message)

    def decode(cell):
        kind = cell.get("t", "n")
        raw = None
        formula = False
        inline = None
        cached = False
        # 普通单元格只有一个值节点；逐项访问避免每个单元格重复编译路径查询。
        for child in cell:
            tag = child.tag
            if tag.endswith("}v"):
                raw = child.text
                cached = True
            elif tag.endswith("}f"):
                formula = True
            elif tag.endswith("}is"):
                inline = child
        if raw == "":
            raw = None
        if formula:
            missing_cache = raw is None and not (kind == "str" and cached)
            note("部分公式没有缓存值，相关单元格不可作为数值依据" if missing_cache else "存在公式，读取文件内缓存值，未重新计算", missing_cache)
        if kind == "inlineStr":
            return (rich_text(inline) if inline is not None else None), True, formula
        if raw is None:
            return None, True, formula
        if len(raw) > 1024 * 1024:
            raise ValueError("Excel 单元格文字超过解析限额")
        if kind == "s":
            index = int(raw)
            if index < 0 or index >= len(shared):
                raise ValueError("Excel 共享文字索引无效")
            return shared[index], True, formula
        if kind in ["str", "e"]:
            return raw, True, formula
        if kind == "b":
            return "True" if raw == "1" else "False", False, formula
        if kind == "d":
            value = from_ISO8601(raw)
            return value.isoformat() if value is not None else None, False, formula
        number = Decimal(raw)
        if not number.is_finite() or abs(number.adjusted()) > 100:
            raise ValueError("Excel 数字超出可读取精度范围")
        style = int(cell.get("s", "0"))
        date, duration = formats[style] if style < len(formats) else (False, False)
        if date or duration:
            value = from_excel(float(number), epoch, timedelta=duration)
            return value.isoformat() if hasattr(value, "isoformat") else str(value), False, formula
        return format(number, "f"), False, formula

    with archive.open(member) as stream:
        for _, element in etree.iterparse(stream, events=("end",), tag=("{*}row", "{*}mergeCell"), resolve_entities=False, no_network=True, load_dtd=False):
            if element.tag.endswith("}mergeCell"):
                merges.append(element.get("ref", ""))
                if len(merges) > 10000:
                    raise ValueError(name + "：合并单元格数量超过解析限额")
                release(element)
                continue
            number = int(element.get("r", str(physical_number + 1)))
            if number <= physical_number or number > MAX_ROWS + 2000:
                raise ValueError(name + "：工作表原始行号超过解析限额或顺序不正确")
            physical_number = number
            values = []
            text_mask = 0
            has_formula = False
            previous_column = 0
            for cell in element:
                if not cell.tag.endswith("}c"):
                    continue
                budget["cells"] += 1
                if budget["cells"] > MAX_CELLS:
                    raise ValueError("Excel 工作簿超过 2000 万个单元格解析限额")
                coordinate = cell.get("r")
                column = column_index_from_string(coordinate.rstrip("0123456789")) if coordinate else previous_column + 1
                if column <= previous_column or column > 1000:
                    raise ValueError(name + "：工作表列数超过解析限额或顺序不正确")
                previous_column = column
                values.extend([None] * (column - len(values)))
                value, is_text, formula = decode(cell)
                has_formula = has_formula or formula
                values[column - 1] = value
                if is_text and value is not None:
                    text_mask |= 1 << (column - 1)
            release(element)
            if not has_formula and not any(value is not None and value != "" for value in values):
                continue
            if has_formula:
                formula_rows.add(number)
            all_text = all(value is None or value == "" or text_mask & (1 << index) for index, value in enumerate(values))
            if start is None:
                start = number
                first_is_text = all_text and not has_formula and not any(numeric_text(value) for value in values)
            elif number > last_number + 1 and all_text and len(gaps) < 20:
                gaps.append((number, values))
            last_number = number
            width = max(width, len(values))
            budget["rows"] += 1
            if count >= MAX_ROWS + 12 or budget["rows"] > MAX_BOOK_ROWS:
                raise ValueError("Excel 单表超过 20 万行或工作簿超过 100 万行解析限额")
            database.execute("INSERT INTO attachment_rows VALUES (?, ?, ?, ?, ?)", (table_id, count, number, json.dumps(values, ensure_ascii=False, separators=(",", ":")), str(text_mask)))
            count += 1
            if count % 2000 == 0:
                database.commit()
    if not count:
        return

    def record(number):
        row = database.execute("SELECT cells, text_mask FROM attachment_rows WHERE table_id = ? AND row_number = ?", (table_id, number)).fetchone()
        values = json.loads(row[0]) if row else []
        return values + [None] * (width - len(values)), int(row[1]) if row else 0

    header_end = start if first_is_text else start - 1
    blocked = []
    ranges = []
    merged_cells = []
    for reference in merges:
        left, top, right, bottom = range_boundaries(reference)
        if bottom > MAX_ROWS + 2000 or right > 1000:
            raise ValueError(name + "：合并单元格范围超过解析限额")
        width = max(width, right)
        values, mask = record(top)
        anchor = values[left - 1]
        is_text = bool(mask & (1 << (left - 1)))
        ranges.append((left, top, right, bottom, anchor, is_text, reference))
        merged_cells.append(dict(range=reference, anchor=get_column_letter(left) + str(top), value=anchor))
        if first_is_text and top == start and bottom > start:
            header_records = list(database.execute("SELECT cells, text_mask FROM attachment_rows WHERE table_id = ? AND row_number <= ? ORDER BY position LIMIT 12", (table_id, bottom))) if bottom - start <= 10 else []
            if header_records and not any(number in formula_rows for number in range(start, bottom + 1)) and all(all(value is None or value == "" or int(mask) & (1 << index) and not numeric_text(value) for index, value in enumerate(json.loads(cells))) for cells, mask in header_records):
                header_end = max(header_end, bottom)
            else:
                blocked.append("合并范围 " + reference + " 跨越表头和数值记录，不能确定数据边界")
    for left, top, right, bottom, anchor, is_text, reference in ranges:
        header_merge = top <= header_end and bottom <= header_end
        label_merge = top > header_end and left == right and is_text
        if header_merge or label_merge:
            for position, cells in database.execute("SELECT position, cells FROM attachment_rows WHERE table_id = ? AND row_number >= ? AND row_number <= ? ORDER BY row_number", (table_id, top, bottom)):
                values = json.loads(cells)
                values.extend([None] * (width - len(values)))
                for index in range(left - 1, right):
                    if values[index] is None:
                        values[index] = anchor
                database.execute("UPDATE attachment_rows SET cells = ? WHERE table_id = ? AND position = ?", (json.dumps(values, ensure_ascii=False, separators=(",", ":")), table_id, position))
            if label_merge:
                local_warnings.append("合并标签 " + reference + " 按原始合并范围展开，原始锚点与范围已保留")
        else:
            blocked.append("合并范围 " + reference + " 含数值或跨列，记录口径有歧义，不能自动汇总")
    header_rows = list(database.execute("SELECT cells FROM attachment_rows WHERE table_id = ? AND row_number <= ? ORDER BY position", (table_id, header_end))) if first_is_text else []
    if header_rows:
        labels = [[] for _ in range(width)]
        for (cells,) in header_rows:
            for index, value in enumerate(json.loads(cells)):
                if value is not None and value.strip() and value.strip() not in labels[index]:
                    labels[index].append(value.strip())
        header = [" / ".join(parts) for parts in labels]
    else:
        header = ["列 " + get_column_letter(index + 1) for index in range(width)]
        blocked.append("首行包含数值或公式，不能确定表头与列含义；全部原始行均保留并使用列标")
    for number, values in gaps:
        if number > header_end and values != header:
            blocked.append("空行后的第 " + str(number) + " 行可能属于另一张表，需确认列含义后再汇总")
    for message in blocked:
        note(message)
    columns = []
    seen = {}
    for index, text in enumerate(header):
        label = text or "列 " + str(index + 1)
        seen[label] = seen.get(label, 0) + 1
        columns.append(label if seen[label] == 1 else label + " [" + str(seen[label]) + "]")
    first_position = len(header_rows)
    row_count = count - first_position
    if row_count > MAX_ROWS:
        raise ValueError(name + "：单表超过 20 万行解析限额")
    location = name + " · A" + str(start) + ":" + get_column_letter(width) + str(last_number)
    item = dict(id=table_id, name=name, columns=columns, rows=[], rowNumbers=[], rowStore=dict(rowCount=row_count, firstPosition=first_position), location=location, warnings=local_warnings)
    if archive_path:
        item["archivePath"] = archive_path
    if merged_cells:
        item["mergedCells"] = merged_cells
    if blocked:
        item["calculationBlocked"] = "；".join(blocked)
    result["tables"].append(item)
    preview = [" | ".join(columns)]
    for (cells,) in database.execute("SELECT cells FROM attachment_rows WHERE table_id = ? AND position >= ? ORDER BY position LIMIT 12", (table_id, first_position)):
        preview.append(" | ".join(value or "" for value in json.loads(cells)))
    block("\n".join(preview), location + " · 表格预览（共 " + str(row_count) + " 条记录）", archive_path, "table")
    if sheet.get("state", "visible") != "visible":
        warning(name + "：该工作表在原文件中处于隐藏状态")
