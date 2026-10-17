"""受限子进程中的确定性文件解析；不读取服务环境和凭据。"""
import csv
import gzip
import bz2
import lzma
import ctypes
import ctypes.util
import io
import json
import os
import pathlib
import re
import resource
import shutil
import subprocess
import sys
import tarfile
import traceback
import zipfile
from html.parser import HTMLParser

INPUT = pathlib.Path(sys.argv[1])
OUT = pathlib.Path(sys.argv[2])
NAME = sys.argv[3]
OUT.mkdir(parents=True, exist_ok=True)
MAX_BYTES = 250 * 1024 * 1024
MAX_CHILD = 100 * 1024 * 1024
MAX_FILES = 100
MAX_ROWS = 200000
MAX_PIXELS = 40000000
result = dict(status="ready", mime="application/octet-stream", blocks=[], tables=[], images=[], entries=[], warnings=[])
budget = dict(bytes=0, files=0, rows=0)

try:
    resource.setrlimit(resource.RLIMIT_CPU, (110, 115))
    resource.setrlimit(resource.RLIMIT_FSIZE, (260 * 1024 * 1024, 260 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_NOFILE, (128, 128))
    if sys.platform != "darwin":
        resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
except (ValueError, OSError):
    raise RuntimeError("无法应用文件解析进程资源限制")


def warning(message):
    if message not in result["warnings"]:
        result["warnings"].append(message)


def error_detail(error):
    # 诊断仅保存在私有任务目录，公共元信息只返回下面的中文摘要。
    with (OUT / "parser-details.txt").open("a", encoding="utf-8") as diagnostic:
        diagnostic.write(traceback.format_exc() + "\n")
    if isinstance(error, ImportError):
        return "文件解析依赖缺失，请运行 npm run install:attachments"
    if isinstance(error, MemoryError):
        return "文件解析超过内存限额，请拆分文件后重试"
    if isinstance(error, csv.Error):
        return "CSV 引号或字段格式不完整，无法可靠解析"
    if type(error).__name__ == "DecompressionBombError":
        return "图片像素超过解析限额"
    detail = str(error).replace(str(INPUT), "上传文件").replace(str(OUT), "解析目录")
    if "password" in detail.lower() or "encrypted" in detail.lower():
        return "文件已加密，当前无法读取"
    if not re.search(r"[\u3400-\u9fff]", detail):
        return "文件损坏或内容格式无法解析"
    return detail[:1000]


def block(text, location, path="", kind="text", image_id=None):
    text = str(text or "").strip()
    if not text and not image_id:
        return
    chunks = [text[i:i + 6000] for i in range(0, max(len(text), 1), 6000)]
    for index, chunk in enumerate(chunks):
        item = dict(id="b" + str(len(result["blocks"]) + 1), kind=kind, text=chunk, location=location + (" · 片段 " + str(index + 1) if len(chunks) > 1 else ""))
        if path:
            item["archivePath"] = path
        if image_id:
            item["imageId"] = image_id
        result["blocks"].append(item)
        if len(result["blocks"]) > 50000:
            raise ValueError("文件结构块数量超过解析限额")


def table(rows, location, path="", name=None, row_numbers=None, table_warnings=None, row_ends=None, merged_cells=None, calculation_blocked=None):
    if not rows:
        return
    width = max(len(row) for row in rows)
    if width > 1000:
        raise ValueError("表格列数超过解析限额")
    header = [str(v).strip() if v is not None else "" for v in rows[0]] + [""] * (width - len(rows[0]))
    seen = {}
    columns = []
    for index, value in enumerate(header):
        label = value or "列 " + str(index + 1)
        seen[label] = seen.get(label, 0) + 1
        columns.append(label if seen[label] == 1 else label + " [" + str(seen[label]) + "]")
    body = [[None if v is None or v == "" else str(v) for v in row] + [None] * (width - len(row)) for row in rows[1:]]
    budget["rows"] += len(body)
    if budget["rows"] > MAX_ROWS:
        raise ValueError("表格记录数超过 200000 条解析限额")
    item = dict(id="t" + str(len(result["tables"]) + 1), name=name or location, columns=columns, rows=body, rowNumbers=row_numbers or list(range(2, len(rows) + 1)), location=location, warnings=table_warnings or [])
    if path:
        item["archivePath"] = path
    if row_ends is not None:
        item["rowEndNumbers"] = row_ends
    if merged_cells:
        item["mergedCells"] = merged_cells
    if calculation_blocked:
        item["calculationBlocked"] = calculation_blocked
    result["tables"].append(item)
    # 预览不是计算输入；全表始终单独保存在 rows 中。
    preview = [" | ".join(columns)] + [" | ".join("" if v is None else v for v in row) for row in body[:12]]
    block("\n".join(preview), location + " · 表格预览（共 " + str(len(body)) + " 条记录）", path, "table")


def decode(data):
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), "UTF-16"
    if b"\x00" in data:
        raise ValueError("无法识别的二进制文件，暂不支持读取")
    for encoding in ["utf-8-sig", "gb18030"]:
        try:
            value = data.decode(encoding)
            if sum(ord(c) < 32 and c not in "\n\r\t" for c in value) > max(2, len(value) // 100):
                continue
            return value, encoding
        except UnicodeDecodeError:
            pass
    raise ValueError("无法识别的文件编码，暂不支持读取")


def numeric_text(value):
    return isinstance(value, str) and re.fullmatch(r"[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?(?:[eE][+-]?\d+)?", value.strip()) is not None


class SafeHTML(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.skip = 0
        self.parts = []
        self.tables = []
        self.current_table = None
        self.row = None
        self.cell = None

    def handle_starttag(self, tag, attrs):
        if tag in ["head", "script", "style", "noscript", "template", "iframe", "object"]:
            self.skip += 1
        if self.skip:
            return
        if tag in ["p", "div", "br", "li", "h1", "h2", "h3", "section", "tr"]:
            self.parts.append("\n")
        if tag == "table":
            self.current_table = []
        if tag == "tr" and self.current_table is not None:
            self.row = []
        if tag in ["th", "td"] and self.row is not None:
            self.cell = []

    def handle_endtag(self, tag):
        if tag in ["head", "script", "style", "noscript", "template", "iframe", "object"]:
            self.skip = max(0, self.skip - 1)
            return
        if self.skip:
            return
        if tag in ["th", "td"] and self.cell is not None:
            self.row.append("".join(self.cell).strip()); self.cell = None
        if tag == "tr" and self.row is not None:
            self.current_table.append(self.row); self.row = None
        if tag == "table" and self.current_table is not None:
            self.tables.append(self.current_table); self.current_table = None
        if tag in ["p", "div", "li", "h1", "h2", "h3", "section"]:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)
            if self.cell is not None:
                self.cell.append(data)


def image(data, location, path=""):
    from PIL import Image
    Image.MAX_IMAGE_PIXELS = MAX_PIXELS
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
    except ImportError:
        pass
    src = Image.open(io.BytesIO(data))
    if src.width * src.height > MAX_PIXELS:
        raise ValueError("图片像素超过解析限额")
    if getattr(src, "n_frames", 1) > 1:
        warning((path or NAME) + "：动画或多帧图片仅读取首帧")
    from PIL import ImageOps
    src = ImageOps.exif_transpose(src).convert("RGB")
    src.thumbnail((2400, 2400))
    image_id = "i" + str(len(result["images"]) + 1)
    src.save(OUT / (image_id + ".jpg"), "JPEG", quality=90)
    item = dict(id=image_id, mime="image/jpeg", width=src.width, height=src.height, location=location)
    if path:
        item["archivePath"] = path
    result["images"].append(item)
    block("图片可供识图，尚未进行模型读取", location, path, "image", image_id)


def safe_zip(path):
    archive = zipfile.ZipFile(path)
    members = archive.infolist()
    if len(members) > 10000 or sum(i.file_size for i in members) > MAX_BYTES:
        archive.close(); raise ValueError("Office 容器展开大小或结构数量超过限额")
    for item in members:
        if item.file_size > MAX_CHILD or item.flag_bits & 1:
            archive.close(); raise ValueError("Office 文件已加密或内部文件超过限额")
    return archive


def convert(src, extension, target):
    executable = shutil.which(os.environ.get("ATTACHMENTS_LIBREOFFICE", "soffice")) or shutil.which("libreoffice")
    if not executable:
        raise ValueError("缺少 LibreOffice 转换器，暂不支持读取该格式")
    if sys.platform == "linux":
        binary = pathlib.Path(executable).resolve().with_name("soffice.bin")
        if binary.is_file():
            executable = str(binary)  # 跳过自行尝试全局 /tmp 管道的桌面启动器。
    folder = OUT / ("conversion-" + str(len(list(OUT.iterdir())))); folder.mkdir()
    original = folder / ("input." + extension); shutil.copyfile(src, original)
    profile = folder / "profile"
    user = profile / "user"; user.mkdir(parents=True)
    (user / "registrymodifications.xcu").write_text('<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item></oor:items>')
    command = [executable, "-env:UserInstallation=" + profile.as_uri(), "-env:OSL_SOCKET_PATH=" + os.environ["TMPDIR"], "-env:UNO_SHARED_PACKAGES_CACHE=" + (folder / "shared-packages").as_uri(), "--headless", "--nologo", "--nodefault", "--nolockcheck", "--convert-to", target, "--outdir", str(folder), str(original)]
    environment = dict(os.environ)
    if sys.platform == "linux":
        shim = next((pathlib.Path(folder) / "attachment_office_socket.so" for folder in sys.path if (pathlib.Path(folder) / "attachment_office_socket.so").is_file()), None)
        if shim is None:
            raise ValueError("缺少受限 Office 转换组件，请运行 npm run install:attachments")
        environment["LD_PRELOAD"] = str(shim)
    for attempt in range(3):
        completed = subprocess.run(command, capture_output=True, timeout=55, env=environment)
        if completed.returncode != 81:  # 初次生成独立配置时，Office 会要求重新启动一次。
            break
    destination = folder / ("input." + target.split(":")[0])
    if completed.returncode or not destination.is_file():
        with (OUT / "parser-details.txt").open("a", encoding="utf-8") as diagnostic:
            diagnostic.write("转换器退出码：" + str(completed.returncode) + "\n" + completed.stdout.decode(errors="replace")[:6000] + "\n" + completed.stderr.decode(errors="replace")[:6000] + "\n")
        raise ValueError("文档转换未完成，请检查文件是否损坏或转换运行环境是否完整")
    return destination


def render_page(src, number, location, path=""):
    executable = shutil.which("pdftoppm")
    if not executable:
        warning((path or NAME) + "：缺少页面渲染器，图形内容未读取")
        return
    destination = OUT / ("page-" + str(len(result["images"]) + 1))
    completed = subprocess.run([executable, "-f", str(number), "-l", str(number), "-singlefile", "-scale-to", "1600", "-png", str(src), str(destination)], capture_output=True, timeout=15)
    png = destination.with_suffix(".png")
    if completed.returncode or not png.is_file():
        warning((path or NAME) + "：页面渲染失败，" + location + " 的图形内容未读取")
        return
    image(png.read_bytes(), location, path); png.unlink()


def parse_docx(src, path):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    archive = safe_zip(src)
    document = Document(src)
    def final_revisions(root):
        changed = False
        for node in root.xpath(".//*[local-name()='del' or local-name()='ins' or local-name()='moveFrom' or local-name()='moveTo']"):
            parent = node.getparent()
            if parent is None:
                continue
            changed = True; kind = node.tag.rsplit("}", 1)[-1]
            if kind == "del" and parent.tag.endswith("}trPr"):
                row = parent.getparent()
                if row is not None and row.getparent() is not None:
                    row.getparent().remove(row)
            elif kind in ["del", "moveFrom"]:
                parent.remove(node)
            else:
                index = parent.index(node)
                for child in list(node):
                    parent.insert(index, child); index += 1
                parent.remove(node)
        if changed:
            warning((path or NAME) + "：含修订，按最终稿保留插入与移入内容，排除删除与移出内容；请按原件核对修订意图")
        return root
    final_revisions(document.element)
    paragraph = 0
    for element in document.element.body:
        if element.tag.endswith("}p"):
            paragraph += 1
            block(Paragraph(element, document).text, "段落 " + str(paragraph), path)
        elif element.tag.endswith("}tbl"):
            tab = Table(element, document)
            table([[cell.text for cell in row.cells] for row in tab.rows], "正文表格 " + str(len(result["tables"]) + 1), path)
            if element.xpath(".//*[local-name()='vMerge' or local-name()='gridSpan' or local-name()='cellDel' or local-name()='cellMerge']"):
                message = "Word 表格含合并单元格或单元格修订，记录粒度需要核对，不能直接自动汇总"
                result["tables"][-1]["calculationBlocked"] = message; result["tables"][-1]["warnings"].append(message); warning((path or NAME) + "：" + message)
    from lxml import etree
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
    for name in archive.namelist():
        if re.match(r"word/(header\d+|footer\d+|footnotes|endnotes)\.xml$", name):
            root = final_revisions(etree.fromstring(archive.read(name), parser))
            texts = root.xpath("//*[local-name()='t']/text()")
            block(" ".join(texts), "附注 " + pathlib.PurePosixPath(name).stem, path)
    root = document.element
    for index, node in enumerate(root.xpath("//*[local-name()='txbxContent']")):
        block(" ".join(node.xpath(".//*[local-name()='t']/text()")), "文本框 " + str(index + 1), path)
    for part in archive.namelist():
        if part.startswith("word/media/"):
            try:
                image(archive.read(part), "内嵌图片 " + pathlib.PurePosixPath(part).name, path)
            except Exception:
                warning((path or NAME) + "：部分内嵌图片无法解码")
    archive.close()


def parse_pptx(src, path):
    from pptx import Presentation
    archive = safe_zip(src); archive.close()
    presentation = Presentation(src)
    visual_pages = []
    def visit(shapes, number):
        visual = False
        for shape in sorted(shapes, key=lambda value: (value.top or 0, value.left or 0)):
            location = "第 " + str(number) + " 页"
            if shape.shape_type == 6:
                visual = visit(shape.shapes, number) or visual
            if shape.has_text_frame:
                block(shape.text, location, path)
            if shape.has_table:
                table([[cell.text for cell in row.cells] for row in shape.table.rows], location + " · 表格", path)
            if shape.has_chart:
                visual = True
                try:
                    categories = [category.label for category in shape.chart.plots[0].categories]
                    rows = [["类别"] + [series.name for series in shape.chart.series]]
                    for index, category in enumerate(categories):
                        rows.append([category] + [series.values[index] for series in shape.chart.series])
                    table(rows, location + " · 图表数据", path)
                except Exception:
                    warning((path or NAME) + "：图表数据不能直接提取，须读取对应页面图像")
            if shape.shape_type == 13:
                visual = True
                try:
                    image(shape.image.blob, location + " · 图片", path)
                except Exception:
                    warning((path or NAME) + "：幻灯片内嵌图无法解码")
            if "graphicFrame" in shape._element.tag and not shape.has_table and not shape.has_chart:
                visual = True
        return visual
    for number, slide in enumerate(presentation.slides, 1):
        before = len(result["blocks"])
        visual = visit(slide.shapes, number)
        if slide.has_notes_slide:
            block(slide.notes_slide.notes_text_frame.text, "第 " + str(number) + " 页 · 备注", path)
        if visual or before == len(result["blocks"]):
            visual_pages.append(number)
    if visual_pages:
        try:
            pdf = convert(src, "pptx", "pdf")
            for number in visual_pages:
                render_page(pdf, number, "第 " + str(number) + " 页 · 整页图像", path)
        except Exception as error:
            warning((path or NAME) + "：幻灯片整体图形尚未覆盖，" + str(error))


def parse_xlsx(src, path):
    with zipfile.ZipFile(src) as container:
        if sum(item.file_size for item in container.infolist()) > 16 * 1024 * 1024:
            from xlsx import parse_workbook
            return parse_workbook(src, OUT, result, block, warning, path)
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries, get_column_letter
    from lxml import etree
    archive = safe_zip(src)
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
    workbook_xml = etree.fromstring(archive.read("xl/workbook.xml"), parser)
    relations = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"), parser)
    targets = {node.get("Id"): node.get("Target", "") for node in relations}
    merged_by_sheet = {}; empty_formula_cache = {}
    for node in workbook_xml.xpath("//*[local-name()='sheet']"):
        target = targets.get(node.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"), "")
        member = target.lstrip("/") if target.startswith("/") else "xl/" + target
        merges = []; empty_cached = set()
        if member in archive.namelist():
            with archive.open(member) as stream:
                for _, element in etree.iterparse(stream, events=("end",), resolve_entities=False, no_network=True, load_dtd=False):
                    if element.tag.endswith("}mergeCell"):
                        merges.append(element.get("ref"))
                    if element.tag.endswith("}c") and element.get("t") == "str" and element.find("{*}f") is not None and element.find("{*}v") is not None:
                        empty_cached.add(element.get("r"))
                    if element.tag.endswith("}c") or element.tag.endswith("}row") or element.tag.endswith("}mergeCell"):
                        element.clear()
                        while element.getprevious() is not None:
                            del element.getparent()[0]
        merged_by_sheet[node.get("name")] = merges
        empty_formula_cache[node.get("name")] = empty_cached
    archive.close()
    value_input = src.open("rb"); formula_input = src.open("rb")
    values = load_workbook(value_input, data_only=True, read_only=True, keep_links=False)
    formulas = load_workbook(formula_input, data_only=False, read_only=True, keep_links=False)
    for sheet in values:
        if sheet.max_row and sheet.max_row > MAX_ROWS + 2000 or sheet.max_column and sheet.max_column > 1000:
            raise ValueError("工作表行列数超过解析限额")
        formula_sheet = formulas[sheet.title]
        records = []; local_warnings = []; blocked = []; formula_rows = set()
        for number, (row, original) in enumerate(zip(sheet.iter_rows(), formula_sheet.iter_rows()), 1):
            output = []
            for cell, formula in zip(row, original):
                value = cell.value
                if formula.data_type == "f":
                    formula_rows.add(number)
                    missing_cache = value is None and formula.coordinate not in empty_formula_cache.get(sheet.title, set())
                    message = sheet.title + "：存在公式，读取文件内缓存值，未重新计算"
                    if missing_cache:
                        message = sheet.title + "：部分公式没有缓存值，相关单元格不可作为数值依据"
                    if missing_cache:
                        warning(message)
                    if message not in local_warnings:
                        local_warnings.append(message)
                if hasattr(value, "isoformat"):
                    value = value.isoformat()
                output.append(value)
            if number not in formula_rows and not any(value is not None and value != "" for value in output):
                continue
            records.append((number, output))
        if records:
            start = records[0][0]; width = max(len(row) for _, row in records)
            matrix = {number: row for number, row in records}
            has_header = start not in formula_rows and all(value is None or isinstance(value, str) for value in records[0][1]) and not any(numeric_text(value) for value in records[0][1])
            header_end = start if has_header else start - 1
            ranges = []; merged_cells = []
            for cell_range in merged_by_sheet.get(sheet.title, []):
                left, top, right, bottom = range_boundaries(cell_range)
                if bottom > MAX_ROWS + 2000 or right > 1000:
                    raise ValueError("合并单元格范围超过解析限额")
                anchor = matrix.get(top, [None] * width)[left - 1]
                ranges.append((left, top, right, bottom, anchor, cell_range))
                merged_cells.append(dict(range=cell_range, anchor=get_column_letter(left) + str(top), value=None if anchor is None else str(anchor)))
                # 只有从首行延伸且没有数值的明确合并表头可合并为多层列名。
                if has_header and top == start and bottom > start:
                    header_records = [row for number, row in records if start <= number <= bottom]
                    if bottom - start <= 10 and not any(number in formula_rows for number in range(start, bottom + 1)) and all(value is None or isinstance(value, str) and not numeric_text(value) for row in header_records for value in row):
                        header_end = max(header_end, bottom)
                    else:
                        blocked.append("合并范围 " + cell_range + " 跨越表头和数值记录，不能确定数据边界")
            for left, top, right, bottom, anchor, cell_range in ranges:
                if top <= header_end and bottom <= header_end:
                    for number, row in records:
                        if top <= number <= bottom:
                            for col in range(left - 1, right):
                                if row[col] is None:
                                    row[col] = anchor
                elif top > header_end and left == right and isinstance(anchor, str):
                    for number, row in records:
                        if top <= number <= bottom and row[left - 1] is None:
                            row[left - 1] = anchor
                    local_warnings.append("合并标签 " + cell_range + " 按原始合并范围展开，原始锚点与范围已保留")
                else:
                    blocked.append("合并范围 " + cell_range + " 含数值或跨列，记录口径有歧义，不能自动汇总")
            if has_header:
                header = []
                for column in range(width):
                    labels = []
                    for number, row in records:
                        value = row[column]
                        if number <= header_end and value is not None and str(value).strip() and str(value).strip() not in labels:
                            labels.append(str(value).strip())
                    header.append(" / ".join(labels))
                body = [(number, row) for number, row in records if number > header_end]
                for index, (number, row) in enumerate(body):
                    if index and number > body[index - 1][0] + 1 and all(value is None or isinstance(value, str) for value in row) and row != header:
                        blocked.append("空行后的第 " + str(number) + " 行可能属于另一张表，原始记录已保留，需确认列含义后再汇总")
            else:
                header = ["列 " + get_column_letter(index + 1) for index in range(width)]
                body = records
                blocked.append("首行包含数值或公式，不能确定表头与列含义；全部原始行均保留并使用列标")
            if blocked:
                local_warnings.extend(blocked)
                for message in blocked:
                    warning(sheet.title + "：" + message)
            location = sheet.title + " · A" + str(start) + ":" + get_column_letter(width) + str(records[-1][0])
            table([header] + [row for _, row in body], location, path, sheet.title, [number for number, _ in body], local_warnings, merged_cells=merged_cells, calculation_blocked="；".join(blocked) if blocked else None)
        if sheet.sheet_state != "visible":
            warning(sheet.title + "：该工作表在原文件中处于隐藏状态")
    values.close(); formulas.close(); value_input.close(); formula_input.close()


def parse_pdf(src, path):
    import pdfplumber
    with pdfplumber.open(src) as document:
        if len(document.pages) > 1000:
            raise ValueError("PDF 页数超过解析限额")
        for number, page in enumerate(document.pages, 1):
            location = "第 " + str(number) + " 页"
            text = page.extract_text() or ""
            block(text, location, path)
            for index, rows in enumerate(page.extract_tables()):
                table(rows, location + " · 表格 " + str(index + 1), path)
            if len(text.strip()) < 40 or page.images:
                render_page(src, number, location + " · 页面图像", path)
                if len(text.strip()) < 40:
                    warning((path or NAME) + "：" + location + " 需要视觉读取，文字抽取不能代表已读懂扫描内容")


def valid_member(name):
    normal = name.replace("\\", "/")
    p = pathlib.PurePosixPath(normal)
    if "\x00" in name or p.is_absolute() or ".." in p.parts or re.match(r"^[A-Za-z]:", normal):
        raise ValueError("压缩包包含越界路径")
    return normal


def archive_items(src, hint):
    if zipfile.is_zipfile(src):
        with zipfile.ZipFile(src) as archive:
            for member in archive.infolist():
                name = valid_member(member.filename)
                mode = member.external_attr >> 16
                if mode & 0o170000 not in [0, 0o100000, 0o040000]:
                    yield name, None, "链接或特殊文件不允许读取"
                    continue
                if member.is_dir():
                    continue
                if member.flag_bits & 1:
                    yield name, None, "文件已加密"
                    continue
                if member.file_size > MAX_CHILD:
                    raise ValueError("压缩包子文件大小超过限额")
                with archive.open(member) as stream:
                    yield name, stream, None
        return
    if tarfile.is_tarfile(src):
        with tarfile.open(src) as archive:
            for member in archive:
                name = valid_member(member.name)
                if member.isdir():
                    continue
                if not member.isfile():
                    yield name, None, "链接或特殊文件不允许读取"
                    continue
                if member.size > MAX_CHILD:
                    raise ValueError("压缩包子文件大小超过限额")
                with archive.extractfile(member) as stream:
                    yield name, stream, None
        return
    prefix = src.open("rb").read(8)
    decoder = gzip.open if prefix.startswith(b"\x1f\x8b") else bz2.open if prefix.startswith(b"BZh") else lzma.open if prefix.startswith(b"\xfd7zXZ") else None
    if decoder:
        suffix = pathlib.Path(hint).suffix
        with decoder(src, "rb") as stream:
            yield pathlib.Path(hint).name[:-len(suffix)] or "解压内容.txt", stream, None
        return
    # RAR / 7Z 通过系统 libarchive 逐项读取，不使用 extractall。
    library = ctypes.util.find_library("archive")
    if not library:
        raise ValueError("缺少 libarchive，暂不支持读取该压缩格式")
    lib = ctypes.CDLL(library)
    pointer = ctypes.c_void_p
    def function(name, restype, argtypes):
        fn = getattr(lib, name); fn.restype = restype; fn.argtypes = argtypes; return fn
    new = function("archive_read_new", pointer, [])
    filters = function("archive_read_support_filter_all", ctypes.c_int, [pointer])
    formats = function("archive_read_support_format_all", ctypes.c_int, [pointer])
    opened = function("archive_read_open_filename", ctypes.c_int, [pointer, ctypes.c_char_p, ctypes.c_size_t])
    next_header = function("archive_read_next_header", ctypes.c_int, [pointer, ctypes.POINTER(pointer)])
    pathname = function("archive_entry_pathname_utf8", ctypes.c_char_p, [pointer])
    filetype = function("archive_entry_filetype", ctypes.c_uint, [pointer])
    symlink = function("archive_entry_symlink", ctypes.c_char_p, [pointer])
    hardlink = function("archive_entry_hardlink", ctypes.c_char_p, [pointer])
    data = function("archive_read_data", ctypes.c_ssize_t, [pointer, pointer, ctypes.c_size_t])
    free = function("archive_read_free", ctypes.c_int, [pointer])
    arc = new(); filters(arc); formats(arc)
    try:
        if opened(arc, os.fsencode(src), 10240) < 0:
            raise ValueError("压缩文件无法打开，可能已加密或损坏")
        entry = pointer()
        while True:
            status = next_header(arc, ctypes.byref(entry))
            if status == 1:
                break
            if status < -20:
                raise ValueError("压缩文件已加密或损坏")
            name = valid_member((pathname(entry) or b"").decode("utf-8", errors="strict"))
            kind = filetype(entry)
            if kind == 0o040000:
                continue
            if kind != 0o100000 or symlink(entry) or hardlink(entry):
                yield name, None, "链接或特殊文件不允许读取"
                continue
            buffer = ctypes.create_string_buffer(65536)
            content = io.BytesIO()
            while True:
                size = data(arc, buffer, len(buffer))
                if size == 0:
                    break
                if size < 0:
                    raise ValueError("压缩子文件已加密或损坏")
                if content.tell() + size > MAX_CHILD or content.tell() + size + budget["bytes"] > MAX_BYTES:
                    raise ValueError("压缩包展开大小超过限额")
                content.write(buffer.raw[:size])
            content.seek(0)
            yield name, content, None
    finally:
        free(arc)


def parse_archive(src, path, depth, hint):
    if depth >= 2:
        raise ValueError("压缩包嵌套深度超过 2 层限额")
    for name, stream, error in archive_items(src, hint):
        budget["files"] += 1
        if budget["files"] > MAX_FILES:
            raise ValueError("压缩包文件数量超过 100 个限额")
        child_path = (path + " / " if path else "") + name
        if error:
            result["entries"].append(dict(path=child_path, status="encrypted" if "加密" in error else "failed", detail=error)); continue
        target = OUT / ("child-" + str(budget["files"]))
        before = len(result["warnings"])
        before_entries = len(result["entries"])
        before_content = (len(result["blocks"]), len(result["tables"]), len(result["images"]))
        try:
            size = 0
            with target.open("wb") as output:
                while True:
                    chunk = stream.read(65536)
                    if not chunk:
                        break
                    size += len(chunk); budget["bytes"] += len(chunk)
                    if size > MAX_CHILD or budget["bytes"] > MAX_BYTES:
                        raise ValueError("压缩包展开大小超过限额")
                    output.write(chunk)
            parse_file(target, name, child_path, depth + 1)
            if before_content == (len(result["blocks"]), len(result["tables"]), len(result["images"])):
                raise ValueError("文件没有可读取正文；可能是仅依赖脚本的页面或空文件")
            incomplete = len(result["warnings"]) > before or any(entry["status"] != "ready" for entry in result["entries"][before_entries:])
            result["entries"].append(dict(path=child_path, status="partial" if incomplete else "ready"))
        except Exception as caught:
            detail = error_detail(caught)
            state = "encrypted" if "加密" in detail else "unsupported" if "不支持" in detail else "failed"
            result["entries"].append(dict(path=child_path, status=state, detail=detail))
        finally:
            target.unlink(missing_ok=True)


def parse_file(src, name, path="", depth=0):
    prefix = src.read_bytes()[:32] if src.stat().st_size < 65536 else src.open("rb").read(32)
    extension = pathlib.Path(name).suffix.lower()
    if not prefix:
        raise ValueError("文件没有可读取内容")
    if prefix.startswith(b"%PDF-"):
        result["mime"] = "application/pdf" if not path else result["mime"]
        return parse_pdf(src, path)
    if zipfile.is_zipfile(src):
        with zipfile.ZipFile(src) as archive:
            members = archive.namelist()
        if "word/document.xml" in members:
            result["mime"] = "application/vnd.openxmlformats-officedocument.wordprocessingml.document" if not path else result["mime"]
            return parse_docx(src, path)
        if "ppt/presentation.xml" in members:
            result["mime"] = "application/vnd.openxmlformats-officedocument.presentationml.presentation" if not path else result["mime"]
            return parse_pptx(src, path)
        if "xl/workbook.xml" in members:
            result["mime"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if not path else result["mime"]
            return parse_xlsx(src, path)
        if "mimetype" in members:
            with zipfile.ZipFile(src) as archive:
                kind = archive.read("mimetype").decode("ascii", errors="ignore")
            target = {"application/vnd.oasis.opendocument.text": ("odt", "docx"), "application/vnd.oasis.opendocument.spreadsheet": ("ods", "xlsx"), "application/vnd.oasis.opendocument.presentation": ("odp", "pptx")}.get(kind)
            if target:
                converted = convert(src, target[0], target[1])
                return parse_file(converted, "input." + target[1], path, depth)
        result["mime"] = "application/zip" if not path else result["mime"]
        return parse_archive(src, path, depth, name)
    if prefix.startswith(b"\xd0\xcf\x11\xe0") or extension in [".rtf", ".odt", ".ods", ".odp"]:
        mapping = {".doc": "docx", ".rtf": "docx", ".odt": "docx", ".ppt": "pptx", ".odp": "pptx", ".xls": "xlsx", ".ods": "xlsx"}
        target = mapping.get(extension)
        if not target:
            raise ValueError("旧 Office 文件类型不明确，暂不支持读取")
        converted = convert(src, extension.lstrip("."), target)
        return parse_file(converted, "input." + target, path, depth)
    if prefix.startswith((b"7z\xbc\xaf\x27\x1c", b"Rar!", b"\x1f\x8b", b"BZh", b"\xfd7zXZ")) or extension in [".tar", ".tgz", ".tbz2"]:
        result["mime"] = "application/x-archive" if not path else result["mime"]
        return parse_archive(src, path, depth, name)
    if prefix.startswith((b"\x89PNG", b"\xff\xd8\xff", b"GIF8", b"BM", b"II*\0", b"MM\0*")) or prefix[8:12] == b"WEBP" or extension in [".heic", ".heif"]:
        result["mime"] = "image/*" if not path else result["mime"]
        return image(src.read_bytes(), "图片", path)
    content, encoding = decode(src.read_bytes())
    if extension in [".csv", ".tsv"]:
        sample = content[:20000]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        except csv.Error:
            dialect = csv.excel_tab if extension == ".tsv" else csv.excel
        reader = csv.reader(io.StringIO(content, newline=""), dialect, strict=True)
        rows = []; numbers = []; starts = []; previous = 0
        for row in reader:
            starts.append(previous + 1); numbers.append(reader.line_num); previous = reader.line_num; rows.append(row)
            if len(rows) > MAX_ROWS + 1:
                raise ValueError("表格记录数超过 200000 条解析限额")
        if rows and any(re.fullmatch(r"[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?", cell.strip()) for cell in rows[0]):
            message = "首行包含数值，不能确定表头与列含义；已保留全部原始行，请补充明确表头后再汇总"
            warning((path or name) + "：" + message)
            header = ["列 " + str(index + 1) for index in range(max(len(row) for row in rows))]
            table([header] + rows, "CSV 表格（" + encoding + "）", path, name, starts, [message], row_ends=numbers, calculation_blocked=message)
        else:
            table(rows, "CSV 表格（" + encoding + "）", path, name, starts[1:], row_ends=numbers[1:])
        result["mime"] = "text/csv" if not path else result["mime"]
    elif extension in [".html", ".htm", ".svg"] or re.search(r"<\s*(?:html|!doctype html|body)\b", content[:1000], re.I):
        if extension == ".svg":
            warning((path or NAME) + "：SVG 仅提取静态文字，未执行或渲染脚本")
        parser = SafeHTML(); parser.feed(content); parser.close()
        paragraphs = [line.strip() for line in "".join(parser.parts).splitlines() if line.strip()]
        for number, paragraph in enumerate(paragraphs, 1):
            block(paragraph, "段落 " + str(number), path)
        for index, rows in enumerate(parser.tables, 1):
            table(rows, "HTML 表格 " + str(index), path)
        result["mime"] = "text/html" if not path else result["mime"]
    else:
        if "<!ENTITY" in content.upper() or "<!DOCTYPE" in content.upper():
            raise ValueError("XML 包含不允许的实体或外部定义")
        for start in range(0, len(content.splitlines()), 60):
            lines = content.splitlines()[start:start + 60]
            block("\n".join(lines), "第 " + str(start + 1) + "—" + str(start + len(lines)) + " 行", path)
        result["mime"] = "text/plain" if not path else result["mime"]


try:
    for suffix in ["", "-journal", "-wal", "-shm"]:
        (OUT / ("tables.sqlite" + suffix)).unlink(missing_ok=True)
    # 全局关闭 XML 外部实体，第三方解析器也不能据上传内容联网。
    from lxml import etree
    etree.set_default_parser(etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False))
    parse_file(INPUT, NAME)
except Exception as caught:
    detail = error_detail(caught)
    warning(detail)
    result["status"] = "encrypted" if "加密" in detail or "password" in detail.lower() else "unsupported" if "不支持" in detail else "failed"

if result["status"] == "ready":
    if not result["blocks"] and not result["tables"] and not result["images"]:
        if result["entries"] and all(entry["status"] == "encrypted" for entry in result["entries"]):
            result["status"] = "encrypted"; warning("压缩包中的文件均已加密，当前无法读取")
        elif result["entries"] and all(entry["status"] == "unsupported" for entry in result["entries"]):
            result["status"] = "unsupported"; warning("压缩包中的文件暂不支持读取")
        else:
            result["status"] = "failed"; warning("文件没有可读取正文；可能是仅依赖脚本的页面或空文件")
    elif result["warnings"] or any(entry["status"] != "ready" for entry in result["entries"]):
        result["status"] = "partial"
elif result["blocks"] or result["tables"] or result["images"]:
    result["status"] = "partial"

(OUT / "parsed.json").write_text(json.dumps(result, ensure_ascii=False, allow_nan=False), encoding="utf-8")
print(json.dumps(dict(status=result["status"], blocks=len(result["blocks"]), tables=len(result["tables"]), images=len(result["images"]))))
