"""合成大工作簿；不包含客户数据。"""
import pathlib
import re
import sys
import zipfile
from xml.etree import ElementTree as XML
from datetime import datetime
from openpyxl import Workbook

root = pathlib.Path(sys.argv[1])
root.mkdir(parents=True, exist_ok=True)
book = Workbook(write_only=True)
for name in ["销售", "触达"]:
    sheet = book.create_sheet(name)
    sheet.append(["编号", "金额", "地区"])
    for number in range(1, 110001):
        sheet.append([str(number).zfill(6), 0.1 if number < 110000 else None, "华东"])
book.save(root / "大表.xlsx")

book = Workbook()
sheet = book.active
sheet.title = "合并标签"
sheet.append(["地区", "金额", "日期", "公式"])
sheet.append(["华东", 0.1, datetime(2026, 10, 16), "=B2"])
sheet.append([None, 0.2, None, None])
sheet.merge_cells("A2:A3")
sheet = book.create_sheet("歧义金额")
sheet.append(["地区", "金额"])
sheet.append(["甲", 10])
sheet.append(["乙", None])
sheet.merge_cells("B2:B3")
sheet = book.create_sheet("跨空列表头")
sheet.append(["销售"])
sheet.merge_cells("A1:D1")
sheet.append(["甲", 10])
sheet = book.create_sheet("省略行号")
sheet.append(["门店", "金额"])
sheet.append([None, None])
sheet.append(["乙", 20])
sheet = book.create_sheet("文字转义")
sheet.append(["编号"])
sheet.append(["共享占位"])
sheet.append(["_x005F_x0041_"])
book.save(root / "语义.xlsx")
with zipfile.ZipFile(root / "语义.xlsx") as archive:
    semantic_content = {name: archive.read(name) for name in archive.namelist()}
semantic_content["xl/worksheets/sheet4.xml"] = re.sub(rb'(<row) r="[0-9]+"', rb'\1', semantic_content["xl/worksheets/sheet4.xml"])
namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
sheet = XML.fromstring(semantic_content["xl/worksheets/sheet5.xml"])
cell = sheet.find(".//{" + namespace + "}c[@r='A2']")
cell.set("t", "s")
for child in list(cell):
    cell.remove(child)
XML.SubElement(cell, "{" + namespace + "}v").text = "0"
semantic_content["xl/worksheets/sheet5.xml"] = XML.tostring(sheet)
semantic_content["xl/sharedStrings.xml"] = ('<sst xmlns="' + namespace + '" count="1" uniqueCount="1"><si><t>_x005F_x000D_</t></si></sst>').encode()
types = XML.fromstring(semantic_content["[Content_Types].xml"])
XML.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", PartName="/xl/sharedStrings.xml", ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml")
semantic_content["[Content_Types].xml"] = XML.tostring(types)
with zipfile.ZipFile(root / "语义.xlsx", "w", zipfile.ZIP_DEFLATED) as archive:
    for name, data in semantic_content.items():
        archive.writestr(name, data)
with zipfile.ZipFile(root / "语义.xlsx", "a", zipfile.ZIP_DEFLATED) as archive:
    archive.writestr("xl/padding.xml", b" " * (17 * 1024 * 1024))

with zipfile.ZipFile(root / "语义.xlsx") as archive:
    content = {name: archive.read(name) for name in archive.namelist() if name != "xl/padding.xml"}
with zipfile.ZipFile(root / "超大内部项.xlsx", "w", zipfile.ZIP_DEFLATED) as archive:
    for name, data in content.items():
        archive.writestr(name, data)
    with archive.open("xl/padding.xml", "w") as stream:
        for _ in range(257):
            stream.write(b" " * (1024 * 1024))
