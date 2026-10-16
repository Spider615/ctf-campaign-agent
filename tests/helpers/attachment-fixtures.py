import csv
import ctypes
import ctypes.util
import bz2
import gzip
import io
import json
import lzma
import os
import pathlib
import subprocess
import sys
import tarfile
import struct
import zlib
import zipfile
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook
from PIL import Image, ImageDraw

root = pathlib.Path(sys.argv[1])
root.mkdir(parents=True, exist_ok=True)
(root / "销售.csv").write_text('门店,实收,说明\n001,0.1,"会员,到店"\n002,0.2,"含换行\n内容"\n003,,缺失\n', encoding="utf-8")
(root / "中文.csv").write_bytes("门店,金额\n东区,30\n".encode("gb18030"))
(root / "小数逗号.csv").write_text("门店;金额\n甲;12,5\n乙;2,5\n", encoding="utf-8")
(root / "无表头.csv").write_text("100\n200\n300\n", encoding="utf-8")
(root / "无表头.tsv").write_text("甲\t100\n乙\t200\n丙\t300\n", encoding="utf-8")
(root / "方案.html").write_text('<html><head><title>会员活动</title><script>偷读密钥</script></head><body><h1>到店方案</h1><p>会员到店礼</p><table><tr><th>项目</th><th>预算</th></tr><tr><td>物料</td><td>300</td></tr></table><script>alert(1)</script></body></html>', encoding="utf-8")
d = Document(); d.add_heading("会员活动", 0); d.add_paragraph("会员到店礼，预算300元。")
t = d.add_table(rows=2, cols=2); t.cell(0, 0).text = "门店"; t.cell(0, 1).text = "预算"; t.cell(1, 0).text = "东区"; t.cell(1, 1).text = "300"
d.save(root / "方案.docx")
from docx.oxml import OxmlElement
d = Document(); paragraph = d.add_paragraph("最终预算")
for tag, text_tag, content in [("w:del", "w:delText", "900元"), ("w:ins", "w:t", "300元")]:
    change = OxmlElement(tag); run = OxmlElement("w:r"); text = OxmlElement(text_tag); text.text = content; run.append(text); change.append(run); paragraph._p.append(change)
d.save(root / "修订.docx")
p = Presentation(); slide = p.slides.add_slide(p.slide_layouts[1]); slide.shapes.title.text = "会员到店活动"; slide.placeholders[1].text = "预算300元，先做门店试点。"; slide.notes_slide.notes_text_frame.text = "备注：需要核验库存"
p.save(root / "活动.pptx")
w = Workbook(); sheet = w.active; sheet.title = "华东"; sheet.append(["门店", "实收", "公式"]); sheet.append(["001", 10.25, "=SUM(B2:B3)"]); sheet.append(["002", 20.75, None]); other = w.create_sheet("华南"); other.append(["门店", "实收"]); other.append(["003", 40]); w.save(root / "销售.xlsx")
w = Workbook(); sheet = w.active; sheet.title = "空行"; sheet.append(["门店", "金额"]); sheet.append(["甲", 10]); sheet.append([None, None]); sheet.append(["乙", 20]); sheet.append(["丙", 30]); w.save(root / "空行.xlsx")
w = Workbook(); sheet = w.active; sheet.title = "合并"; sheet.append(["地区", "金额"]); sheet.append(["华东", 10]); sheet.append([None, 20]); sheet.merge_cells("A2:A3"); w.save(root / "合并.xlsx")
w = Workbook(); sheet = w.active; sheet.title = "多层表头"; sheet.append(["地区", "金额", None]); sheet.append([None, "实收", "成本"]); sheet.merge_cells("A1:A2"); sheet.merge_cells("B1:C1"); sheet.append(["华东", 30, 20]); w.save(root / "表头.xlsx")
w = Workbook(); sheet = w.active; sheet.append(["门店", "金额"]); sheet.append(["甲", 10]); sheet.append(["乙", None]); sheet.merge_cells("B2:B3"); w.save(root / "合并金额.xlsx")
image = Image.new("RGB", (320, 180), "white"); ImageDraw.Draw(image).text((20, 30), "CTF 300", fill="black"); image.save(root / "物料.png")
large = bytearray((root / "物料.png").read_bytes()); large[16:24] = struct.pack(">II", 9000, 9000); large[29:33] = struct.pack(">I", zlib.crc32(large[12:29])); (root / "超像素.png").write_bytes(large)
(root / "实体.xml").write_text('<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><x>&e;</x>', encoding="utf-8")
(root / "损坏.docx").write_bytes(b"PK\x03\x04broken\x00")
image.save(root / "扫描.pdf", "PDF")
import pillow_heif
pillow_heif.register_heif_opener()
image.save(root / "物料.heic", "HEIF")
for suffix, compressor in [("gz", gzip.compress), ("bz2", bz2.compress), ("xz", lzma.compress)]:
    (root / ("销售.csv." + suffix)).write_bytes(compressor((root / "销售.csv").read_bytes()))
with tarfile.open(root / "资料.tar.gz", "w:gz") as archive:
    archive.add(root / "销售.csv", arcname="资料/销售.csv")
libarchive = ctypes.CDLL(ctypes.util.find_library("archive"))
def archive_function(name, output_type, input_types):
    fn = getattr(libarchive, name); fn.restype = output_type; fn.argtypes = input_types; return fn
pointer = ctypes.c_void_p
new_archive = archive_function("archive_write_new", pointer, [])
set_7z = archive_function("archive_write_set_format_7zip", ctypes.c_int, [pointer])
open_archive = archive_function("archive_write_open_filename", ctypes.c_int, [pointer, ctypes.c_char_p])
new_entry = archive_function("archive_entry_new", pointer, [])
set_path = archive_function("archive_entry_set_pathname", None, [pointer, ctypes.c_char_p])
set_size = archive_function("archive_entry_set_size", None, [pointer, ctypes.c_int64])
set_kind = archive_function("archive_entry_set_filetype", None, [pointer, ctypes.c_uint])
set_perm = archive_function("archive_entry_set_perm", None, [pointer, ctypes.c_uint])
write_header = archive_function("archive_write_header", ctypes.c_int, [pointer, pointer])
write_data = archive_function("archive_write_data", ctypes.c_ssize_t, [pointer, ctypes.c_char_p, ctypes.c_size_t])
free_entry = archive_function("archive_entry_free", None, [pointer])
close_archive = archive_function("archive_write_close", ctypes.c_int, [pointer])
free_archive = archive_function("archive_write_free", ctypes.c_int, [pointer])
archive = new_archive(); assert set_7z(archive) == 0; assert open_archive(archive, os.fsencode(root / "资料.7z")) == 0
entry = new_entry(); content = (root / "销售.csv").read_bytes(); set_path(entry, b"data.csv"); set_size(entry, len(content)); set_kind(entry, 0o100000); set_perm(entry, 0o644)
assert write_header(archive, entry) == 0; assert write_data(archive, content, len(content)) == len(content)
free_entry(entry); assert close_archive(archive) == 0; free_archive(archive)
(root / "空页面.html").write_text('<html><head><title>尚未加载</title></head><body><div id="root"></div><script>document.querySelector("#root").textContent="不能执行"</script></body></html>', encoding="utf-8")
with zipfile.ZipFile(root / "混合.zip", "w", zipfile.ZIP_DEFLATED) as z:
    for name in ["方案.docx", "销售.csv", "物料.png"]: z.write(root / name, "资料/" + name)
    z.writestr("未知.bin", bytes([0, 255, 1, 128]))
with zipfile.ZipFile(root / "越界.zip", "w") as z: z.writestr("../../secret.txt", "不得写出")
with zipfile.ZipFile(root / "链接.zip", "w") as z:
    z.writestr("safe.txt", "普通正文可读")
    link = zipfile.ZipInfo("link.txt"); link.create_system = 3; link.external_attr = 0o120777 << 16; z.writestr(link, "../../private-canary")
encrypted = io.BytesIO()
with zipfile.ZipFile(encrypted, "w") as z: z.writestr("locked.txt", "不能作为已读取来源")
marked = bytearray(encrypted.getvalue())
for signature, flag_offset in [(b"PK\x03\x04", 6), (b"PK\x01\x02", 8)]:
    index = marked.index(signature) + flag_offset; marked[index:index + 2] = struct.pack("<H", struct.unpack("<H", marked[index:index + 2])[0] | 1)
(root / "加密标记.zip").write_bytes(marked)
with zipfile.ZipFile(root / "超限.zip", "w", zipfile.ZIP_DEFLATED) as z:
    for n in range(102): z.writestr(str(n) + ".txt", "文件")
if "--legacy" in sys.argv:
    os.environ["PYTHONDONTWRITEBYTECODE"] = "1"
    for filename, target in [("方案.docx", "doc"), ("活动.pptx", "ppt"), ("销售.xlsx", "xls")]:
        converted = subprocess.run([os.environ.get("ATTACHMENTS_LIBREOFFICE", "soffice"), "-env:UserInstallation=" + (root / "fixture-office").as_uri(), "--headless", "--convert-to", target, "--outdir", str(root), str(root / filename)], capture_output=True, timeout=45)
        if converted.returncode or not (root / pathlib.Path(filename).with_suffix("." + target)).is_file():
            raise RuntimeError("合成旧格式文件转换失败：" + converted.stderr.decode(errors="replace"))
print(json.dumps({"files": len(list(root.iterdir()))}))
